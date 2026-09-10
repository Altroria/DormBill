"""检查Excel文件结构"""
import openpyxl
from pathlib import Path

# 读取Excel文件
excel_path = Path(r"e:\cursor\DormBill\7月宿舍员工 扣款.xlsx")
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb.active

# 读取表头
headers = []
for cell in ws[1]:
    headers.append(cell.value)

print("表头信息：")
for idx, header in enumerate(headers):
    print(f"列{idx}: {header}")

print("\n前5行数据示例：")
for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=6, values_only=True), start=2):
    print(f"\n第{row_idx}行:")
    for idx, value in enumerate(row):
        print(f"  列{idx} ({headers[idx] if idx < len(headers) else ''}): {value}")
