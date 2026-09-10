"""Excel 导入导出服务"""
from io import BytesIO
from datetime import date
from decimal import Decimal
from typing import List, Dict
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from ..utils.decimal_utils import to_decimal


# 样式
HEADER_FONT = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
MONEY_ALIGN = Alignment(horizontal="right", vertical="center")
BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)


def _style_header(ws, row=1):
    """设置表头样式"""
    for cell in ws[row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER


def _autosize(ws):
    """自动列宽"""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    length = max(len(str(c) for c in str(cell.value).split("\n")))
                    max_len = max(max_len, length)
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 30)


def _parse_date_cell(cell) -> date | None:
    """解析日期单元格"""
    if cell is None or cell == "":
        return None
    if isinstance(cell, date):
        return cell
    if isinstance(cell, str):
        try:
            return date.fromisoformat(cell.replace("/", "-"))
        except Exception:
            return None
    return None


# ============ 导出 ============

def export_settlement_excel(items: List[Dict]) -> bytes:
    """导出员工扣款表"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "员工扣款表"

    headers = [
        "楼号", "房号", "室号", "房间", "任职单位", "一级部门", "职务",
        "工号", "姓名", "转宿日期，备注",
        "普通电费", "空调电费", "水费",
        "应住房租", "实扣房租", "补扣-", "补加+", "水电+实扣+补扣",
    ]
    ws.append(headers)
    _style_header(ws)

    for item in items:
        ws.append([
            item.get("building_no", ""),
            item.get("room_no", ""),
            item.get("room_unit", ""),
            item.get("room_name", ""),
            item.get("company", ""),
            item.get("department", ""),
            item.get("position", ""),
            item.get("employee_no", ""),
            item.get("employee_name", ""),
            item.get("remark", ""),
            item.get("electricity_fee", 0),
            item.get("ac_electricity_fee", 0),
            item.get("water_fee", 0),
            item.get("rent_should", 0),
            item.get("rent_actual", 0),
            item.get("deduction_minus", 0),
            item.get("deduction_plus", 0),
            item.get("total_amount", 0),
        ])

    # 数字列右对齐
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=11, max_col=18):
        for cell in row:
            cell.alignment = MONEY_ALIGN
            cell.border = BORDER

    # 数字格式
    for col in range(11, 19):
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col).number_format = "0.00"

    _autosize(ws)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.getvalue()


def export_meter_detail_excel(items: List[Dict]) -> bytes:
    """导出房间电费明细"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "房间电费明细"

    headers = [
        "楼栋", "房号", "房间", "电表编号",
        "上月普通读数", "本月普通读数", "普通用电量", "电价", "普通电费",
        "上月空调读数", "本月空调读数", "空调度数差", "空调平均单价", "空调电费",
    ]
    ws.append(headers)
    _style_header(ws)

    for item in items:
        ws.append([
            item.get("building_no", ""),
            item.get("room_no", ""),
            item.get("room_name", ""),
            item.get("meter_no", ""),
            item.get("previous_reading", 0),
            item.get("current_reading", 0),
            item.get("total_degree", 0),
            item.get("electricity_price", 0),
            item.get("total_fee", 0),
            item.get("ac_previous_reading", 0),
            item.get("ac_current_reading", 0),
            item.get("ac_degree", 0),
            item.get("ac_unit_price", 0),
            item.get("ac_fee", 0),
        ])

    for col in range(5, 15):
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            cell.alignment = MONEY_ALIGN
            cell.border = BORDER
            cell.number_format = "0.00"

    _autosize(ws)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.getvalue()


def export_water_detail_excel(items: List[Dict]) -> bytes:
    """导出水费分摊明细"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "水费分摊明细"

    headers = [
        "水费周期起始", "水费周期结束", "楼栋", "房号", "房间",
        "姓名", "第一个月天数", "第二个月天数", "是否有效", "个人水费", "备注",
    ]
    ws.append(headers)
    _style_header(ws)

    for item in items:
        ws.append([
            item.get("period_start", ""),
            item.get("period_end", ""),
            item.get("building_no", ""),
            item.get("room_no", ""),
            item.get("room_name", ""),
            item.get("employee_name", ""),
            item.get("month1_days", 0),
            item.get("month2_days", 0),
            "有效" if item.get("is_valid") else "无效",
            item.get("amount", 0),
            item.get("remark", ""),
        ])

    for col in [7, 8, 10]:
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            cell.alignment = MONEY_ALIGN
            cell.border = BORDER
            cell.number_format = "0.00"

    _autosize(ws)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.getvalue()


# ============ 导入 ============

def import_employees_from_excel(bio: BytesIO) -> Dict:
    """
    导入员工

    期望列：工号 | 姓名 | 任职单位 | 一级部门 | 职务
    """
    wb = openpyxl.load_workbook(bio, data_only=True)
    ws = wb.active
    rows = []
    errors = []
    headers = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(c or "").strip() for c in row]
            continue
        if not any(row):
            continue
        try:
            d = dict(zip(headers, row))
            emp_no = str(d.get("工号") or "").strip()
            name = str(d.get("姓名") or "").strip()
            if not emp_no or not name:
                errors.append({"row": i + 1, "msg": "工号或姓名为空"})
                continue
            rows.append({
                "employee_no": emp_no,
                "name": name,
                "company": str(d.get("任职单位") or "").strip() or None,
                "department": str(d.get("一级部门") or "").strip() or None,
                "position": str(d.get("职务") or "").strip() or None,
                "status": "active",
            })
        except Exception as e:
            errors.append({"row": i + 1, "msg": str(e)})
    return {"rows": rows, "errors": errors}


def import_rooms_from_excel(bio: BytesIO) -> Dict:
    """
    导入房间

    期望列：楼栋 | 房号 | 房间 | 房租标准 | 电表编号 | 默认电价
    """
    wb = openpyxl.load_workbook(bio, data_only=True)
    ws = wb.active
    rows = []
    errors = []
    headers = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(c or "").strip() for c in row]
            continue
        if not any(row):
            continue
        try:
            d = dict(zip(headers, row))
            bno = str(d.get("楼栋") or "").strip()
            rno = str(d.get("房号") or "").strip()
            rname = str(d.get("房间") or "").strip()
            if not bno or not rno or not rname:
                errors.append({"row": i + 1, "msg": "楼栋/房号/房间不能为空"})
                continue
            rent = to_decimal(d.get("房租标准", 0))
            price = to_decimal(d.get("默认电价", 0.49))
            rows.append({
                "building_no": bno,
                "room_no": rno,
                "room_name": rname,
                "rent_standard": rent,
                "meter_no": str(d.get("电表编号") or "").strip() or None,
                "electricity_price": price,
                "status": "active",
            })
        except Exception as e:
            errors.append({"row": i + 1, "msg": str(e)})
    return {"rows": rows, "errors": errors}


def import_residences_from_excel(bio: BytesIO) -> Dict:
    """
    导入入住记录

    期望列：工号/姓名 | 楼栋 | 房号 | 房间 | 入住日期 | 试用期月数 | 备注
    """
    wb = openpyxl.load_workbook(bio, data_only=True)
    ws = wb.active
    rows = []
    errors = []
    headers = None
    from ..database import SessionLocal
    db = SessionLocal()
    try:
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(c or "").strip() for c in row]
                continue
            if not any(row):
                continue
            try:
                d = dict(zip(headers, row))
                emp_key = str(d.get("工号/姓名") or d.get("工号") or "").strip()
                bno = str(d.get("楼栋") or "").strip()
                rno = str(d.get("房号") or "").strip()
                rname = str(d.get("房间") or "").strip()
                check_in = _parse_date_cell(d.get("入住日期"))

                from ..models import Building, Room, Employee
                emp = db.query(Employee).filter(
                    (Employee.employee_no == emp_key) | (Employee.name == emp_key)
                ).filter(Employee.deleted_at.is_(None)).first()
                if not emp:
                    errors.append({"row": i + 1, "msg": f"员工不存在: {emp_key}"})
                    continue
                building = db.query(Building).filter(Building.building_no == bno).first()
                if not building:
                    errors.append({"row": i + 1, "msg": f"楼栋不存在: {bno}"})
                    continue
                room = db.query(Room).filter(
                    Room.building_id == building.id,
                    Room.room_no == rno,
                    Room.room_name == rname,
                ).first()
                if not room:
                    errors.append({"row": i + 1, "msg": f"房间不存在: {bno}-{rno}-{rname}"})
                    continue

                rows.append({
                    "employee_id": emp.id,
                    "room_id": room.id,
                    "check_in_date": check_in or date.today(),
                    "is_primary_payer": 0,
                    "probation_months": int(d.get("试用期月数") or 0),
                    "status": "valid",
                    "remark": str(d.get("备注") or "").strip() or None,
                })
            except Exception as e:
                errors.append({"row": i + 1, "msg": str(e)})
    finally:
        db.close()
    return {"rows": rows, "errors": errors}


__all__ = [
    "export_settlement_excel",
    "export_meter_detail_excel",
    "export_water_detail_excel",
    "import_employees_from_excel",
    "import_rooms_from_excel",
    "import_residences_from_excel",
]
