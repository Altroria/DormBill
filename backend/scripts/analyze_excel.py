"""
查看Excel文件详细列结构
"""
import openpyxl
from pathlib import Path

def analyze_excel_structure(file_path, file_name):
    """分析Excel文件结构"""
    print(f"\n{'='*60}")
    print(f"分析文件: {file_name}")
    print('='*60)
    
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    
    # 打印表头
    print("\n列名 (第2行):")
    headers = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
    for idx, header in enumerate(headers):
        if header:
            print(f"  列{idx}: {header}")
    
    # 打印前5行数据
    print("\n前5行数据:")
    for row_idx in range(3, 8):
        print(f"\n第{row_idx}行:")
        row_data = list(ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))[0]
        for col_idx, value in enumerate(row_data):
            if value and col_idx < 10:  # 只看前10列
                print(f"  列{col_idx} ({headers[col_idx] if col_idx < len(headers) else ''}): {value}")

def main():
    base_dir = Path(__file__).parent.parent.parent
    
    # 分析员工扣款文件
    employee_file = base_dir / "7月宿舍员工 扣款.xlsx"
    if employee_file.exists():
        analyze_excel_structure(employee_file, "员工扣款文件")
    
    # 分析空调分摊文件
    ac_file = base_dir / "7月空调各房间分摊费用.xlsx"
    if ac_file.exists():
        analyze_excel_structure(ac_file, "空调分摊文件")

if __name__ == "__main__":
    main()
