"""
导入7月份数据脚本 - 第一步: 预览数据
"""
import openpyxl
from pathlib import Path

def read_employee_deduction_file(file_path):
    """读取员工扣款文件"""
    print(f"\n正在读取文件: {file_path}")
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    
    print(f"工作表名称: {ws.title}")
    print(f"总行数: {ws.max_row}, 总列数: {ws.max_column}")
    
    # 打印前10行数据
    print("\n前10行数据预览:")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        print(f"第{i}行: {row}")
    
    return wb, ws

def read_ac_allocation_file(file_path):
    """读取空调分摊费用文件"""
    print(f"\n正在读取文件: {file_path}")
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    
    print(f"工作表名称: {ws.title}")
    print(f"总行数: {ws.max_row}, 总列数: {ws.max_column}")
    
    # 打印前10行数据
    print("\n前10行数据预览:")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        print(f"第{i}行: {row}")
    
    return wb, ws

def main():
    # 文件路径
    base_dir = Path(__file__).parent.parent.parent
    employee_file = base_dir / "7月宿舍员工 扣款.xlsx"
    ac_file = base_dir / "7月空调各房间分摊费用.xlsx"
    
    print("=" * 60)
    print("开始导入7月份数据")
    print("=" * 60)
    
    # 检查文件是否存在
    if not employee_file.exists():
        print(f"错误: 找不到文件 {employee_file}")
        return
    
    if not ac_file.exists():
        print(f"错误: 找不到文件 {ac_file}")
        return
    
    # 读取文件
    emp_wb, emp_ws = read_employee_deduction_file(employee_file)
    ac_wb, ac_ws = read_ac_allocation_file(ac_file)
    
    print("\n文件读取成功!")
    print("请查看上面的数据预览,确认数据格式正确")

if __name__ == "__main__":
    main()
